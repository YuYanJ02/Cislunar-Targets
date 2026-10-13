#!/usr/bin/env python3
"""Import the two curated Excel registries into the web catalog schema."""

from __future__ import annotations

import argparse
import json
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def iso_value(value: Any) -> Any:
    if isinstance(value, datetime):
        if value.tzinfo is None:
            value = value.replace(tzinfo=timezone.utc)
        return value.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")
    if isinstance(value, date):
        return value.isoformat()
    return value


def rows_from_xlsx(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
    rows: list[dict[str, Any]] = []
    for values in iterator:
        if not any(value is not None for value in values):
            continue
        rows.append({header: iso_value(value) for header, value in zip(headers, values) if header})
    return rows


def source_from_row(row: dict[str, Any], source: str) -> dict[str, Any]:
    orbit = {
        "epoch_utc": row.get("tle_epoch_utc"),
        "semi_major_axis_km": row.get("semi_major_axis_km"),
        "period_days": row.get("period_days"),
        "ephemeris_type": str(row.get("ephemeris_type")) if row.get("ephemeris_type") is not None else None,
        "semi_major_axis_method": row.get("semi_major_axis_method"),
        "eccentricity": None,
        "inclination_deg": None,
        "raan_deg": None,
        "arg_pericenter_deg": None,
        "mean_anomaly_deg": None,
        "mean_motion_rev_day": None,
        "perigee_altitude_km": None,
        "apogee_altitude_km": None,
        "tle_line0": None,
        "tle_line1": None,
        "tle_line2": None,
    }
    return {
        "source": source,
        "eligible": str(row.get("include", "TRUE")).upper() in {"TRUE", "1", "YES"},
        "last_verified_utc": row.get("last_verified_utc"),
        "source_url": row.get("source_url"),
        "source_commit": row.get("source_commit"),
        "bill_gray_file": row.get("bill_gray_file"),
        "bill_gray_name": row.get("bill_gray_name"),
        "catalog_id": row.get("bill_gray_catalog_id") or row.get("norad_cat_id"),
        "orbit": orbit,
    }


def key_for(row: dict[str, Any]) -> str:
    if row.get("norad_cat_id"):
        return f"NORAD:{str(row['norad_cat_id']).strip()}"
    if row.get("cospar_id"):
        return f"COSPAR:{str(row['cospar_id']).strip().upper()}"
    return str(row.get("canonical_object_id") or row.get("preferred_name"))


def split_aliases(value: Any) -> list[str]:
    if not value:
        return []
    return [part.strip() for part in str(value).replace(",", "|").split("|") if part.strip()]


def merge_rows(space_track_rows: list[dict[str, Any]], bill_gray_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    records: dict[str, dict[str, Any]] = {}
    for source, rows in (("spacetrack", space_track_rows), ("bill_gray", bill_gray_rows)):
        for row in rows:
            key = key_for(row)
            record = records.setdefault(key, {
                "key": key,
                "name": row.get("preferred_name") or row.get("bill_gray_name") or key,
                "norad_cat_id": str(row.get("norad_cat_id")) if row.get("norad_cat_id") is not None else None,
                "cospar_id": row.get("cospar_id"),
                "aliases": [],
                "object_type": row.get("object_type") or "spacecraft",
                "orbit_regime": row.get("orbit_regime") or "XGEO",
                "primary_body": row.get("primary_body") or "Earth",
                "identity_status": row.get("identity_status") or "unknown",
                "mission_status": row.get("mission_status") or "unknown",
                "priority": row.get("priority") or "medium",
                "notes": [],
                "eligible": False,
                "sources": [],
            })
            if source == "bill_gray" and row.get("preferred_name"):
                record["name"] = row["preferred_name"]
            record["aliases"] = sorted(set(record["aliases"] + split_aliases(row.get("aliases"))))
            if row.get("notes") and row["notes"] not in record["notes"]:
                record["notes"].append(row["notes"])
            source_entry = source_from_row(row, source)
            record["sources"] = [item for item in record["sources"] if item["source"] != source]
            record["sources"].append(source_entry)
            record["eligible"] = record["eligible"] or source_entry["eligible"]
    return sorted(records.values(), key=lambda item: (item["name"].casefold(), item["key"]))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--bill-gray-xlsx", type=Path, required=True)
    parser.add_argument("--space-track-xlsx", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    records = merge_rows(rows_from_xlsx(args.space_track_xlsx), rows_from_xlsx(args.bill_gray_xlsx))
    now = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    catalog = {
        "schema_version": "1.0.0",
        "generated_at_utc": now,
        "selection": {
            "launch_date_after": "2000-12-31",
            "object_type": "PAYLOAD",
            "decay_date": None,
            "semi_major_axis_min_km": 42164,
        },
        "attribution": {
            "space_track": "U.S. Space Command basic SSA data via Space-Track.org",
            "bill_gray": "Bill Gray high-altitude artificial-satellite TLE repository",
        },
        "records": records,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(catalog, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {len(records)} unified records to {args.output}")


if __name__ == "__main__":
    main()

